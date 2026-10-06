"""Judge hand-check sheet, Phase 5 step 1 (EVAL §5, D-045): every dev premise item, baseline and
treatment, with the judge's decision and blank columns for the owner's own labels. No model calls.

Treatment answers are the baseline answers with the frozen verifier's verdict applied
(`results/verifier/dev_verifier_v2.jsonl`, full mode). An answer the verifier did not block is
the same text, so its judge decision is the baseline one; a blocked answer gets no judge decision
(the scorer decides it under EVAL §3).

Usage: python -m filingguard.eval.judge_sheet
"""

import json
from pathlib import Path

from filingguard.eval.dev_report import is_premise, load_dev, load_judge
from filingguard.eval.judge import INPUTS

VERIFIER_RUN = Path("results/verifier/dev_verifier_v2.jsonl")
SHEET = Path("review/judge_handcheck_dev_v1.xlsx")
FIELDS = ("rejects_premise", "confirms_change", "gives_reasons", "reasons_in_filing", "says_no_reasons_given")
REASONS_CHARS = 2500


def rows() -> list[list]:
    items, gold, runs = load_dev()
    judge = load_judge()
    entries = json.loads(INPUTS.read_text(encoding="utf-8"))["items"]
    verdicts = {r["source_id"]: r for r in map(json.loads, VERIFIER_RUN.read_text(encoding="utf-8").splitlines())
                if r and r["kind"] == "natural" and r["mode"] == "full"}
    out = []
    for it in sorted((i for i in items if is_premise(i)), key=lambda i: i["id"]):
        entry = entries[gold[it["id"]]["spec_id"]]
        reasons = "\n---\n".join(entry["reasons"])[:REASONS_CHARS] if entry["reasons"] else "(none)"
        decision = judge[f"{it['id']}:baseline"]
        v = verdicts[it["id"]]
        for condition in ("baseline", "treatment"):
            blocked = condition == "treatment" and v["decision"] == "block"
            note = ("" if condition == "baseline" else
                    "blocked: no judge decision" if blocked else
                    f"verifier {v['decision']}: same answer as baseline, same judge decision")
            judged = [None] * len(FIELDS) + [None] if blocked else [decision.get(f) for f in FIELDS] + [decision["reason"]]
            out.append([it["id"], it["attack_type"], it["subtype"], condition, note, it["question"], entry["premise"],
                        entry["values"], reasons, runs[it["id"]]["answer"]["text"], *judged,
                        *([""] * len(FIELDS)), ""])
    return out


def write(path: Path = SHEET) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    data = rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "judge hand-check"
    ws.append(["Item", "Type", "Subtype", "Condition", "Treatment note", "Question", "Premise", "Reported values",
               "Reason passages given to the judge", "Answer", *[f"Judge: {f}" for f in FIELDS], "Judge: reason",
               *[f"Your {f}" for f in FIELDS], "Your note"])
    for r in data:
        ws.append(r)
    widths = [9, 13, 18, 10, 26, 40, 34, 26, 60, 60] + [11] * len(FIELDS) + [40] + [11] * len(FIELDS) + [30]
    for n, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, n).column_letter].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "E2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return len(data)


if __name__ == "__main__":
    print(f"wrote {write()} rows to {SHEET}")
