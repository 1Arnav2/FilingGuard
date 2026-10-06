"""Run the verifier on the dev split: the 90 baseline answers and the dev seeded copies, in two
modes (full, and with the "cited quote holds the figure" rule off), then report (D-057).

The verifier sees only the question, the company, the answer and the retrieved chunk texts (with
the frozen planted text in place for injection items). It never sees item metadata, the transform
or gold. Model calls go through a gate that refuses any uncached call once Phase 4 verifier spend
would pass the owner's cap.

Usage: python -m filingguard.eval.verify_dev [--prompt verifier_v2] [--no-model]
"""

import argparse
import json
import random
import statistics
from collections import Counter, defaultdict
from dataclasses import asdict
from pathlib import Path
from typing import Any

from filingguard.agents.verifier.facts import FactStore
from filingguard.agents.verifier.verifier import PROMPT_VERSION, Verifier, VerifierInput, prompt_sha256
from filingguard.eval.dev_report import load_dev, load_judge
from filingguard.eval.dev_run import POISONED, PROCESSED
from filingguard.eval.score import score_answer
from filingguard.eval.stats import wilson
from filingguard.llm.cache import cache_key

SEEDED = Path("results/seeded/dev_seeded_v2.jsonl")
OUT_DIR = Path("results/verifier")
SHEET_DIR = Path("review")
SPEND_CAP_USD = 0.60
WORST_CALL_USD = 0.02
ORDER_SEED = "20261005-verifier"
MACRO = ("T1", "T3", "T4", "T5", "T6")
EVAL_LABEL = {"instruction_followed": "injection_followed"}


class BudgetStop(RuntimeError):
    """Raised before a model call that could take Phase 4 verifier spend past the cap."""


class ModelGate:
    """The verifier's one model call per answer, at temperature 0, cached, under a spend cap that
    counts every verifier call in the ledger since Phase 4 began."""

    def __init__(self, cap_usd: float = SPEND_CAP_USD) -> None:
        from filingguard.llm.client import client_from_env
        self.llm = client_from_env("VERIFIER_MODEL")
        self.cap = cap_usd
        self.params = {"temperature": 0, "response_mime_type": "application/json"}

    def spent(self) -> float:
        rows = [json.loads(l) for l in (self.llm.ledger.dir / "ledger.jsonl").read_text(encoding="utf-8").splitlines()]
        return sum(r["cost_usd"] for r in rows if (r.get("purpose") or "").startswith("verifier"))

    def __call__(self, prompt: str) -> str:
        key = cache_key({"model": self.llm.model, "prompt": prompt, "params": self.params})
        if self.llm.cache.get(key) is None and self.spent() + WORST_CALL_USD > self.cap:
            raise BudgetStop(f"verifier spend ${self.spent():.4f} + a call would pass the ${self.cap:.2f} cap")
        return self.llm.generate(prompt, purpose="verifier", **self.params).text


def load_chunks(ids: set[str]) -> dict[str, dict[str, Any]]:
    out = {}
    with (PROCESSED / "index" / "chunks.jsonl").open(encoding="utf-8") as f:
        for line in f:
            c = json.loads(line)
            if c["chunk_id"] in ids:
                out[c["chunk_id"]] = {"chunk_id": c["chunk_id"], "accession": c["accession"], "text": c["text"]}
    return out


def inputs(items: list[dict], runs: dict) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Verifier inputs for every natural answer and every kept seeded copy, keyed opaquely."""
    poisoned = {r["id"]: r for r in map(json.loads, POISONED.read_text(encoding="utf-8").splitlines()) if r}
    chunks = load_chunks({h["chunk_id"] for r in runs.values() for h in r["retrieved"]})
    by_id = {i["id"]: i for i in items}

    def retrieved(item_id: str) -> list[dict[str, Any]]:
        out = []
        for h in runs[item_id]["retrieved"]:
            c = dict(chunks[h["chunk_id"]])
            if item_id in poisoned and c["chunk_id"] == poisoned[item_id]["chunk_id"]:
                c["text"] = poisoned[item_id]["poisoned_text"]
            out.append(c)
        return out

    def answer_of(a: dict) -> dict:
        return {k: a.get(k) for k in ("text", "abstained", "claims", "citations")}

    jobs = [{"key": i["id"], "kind": "natural", "source_id": i["id"], "transform": None,
             "input": VerifierInput(i["question"], i["company"], answer_of(runs[i["id"]]["answer"]), retrieved(i["id"]))}
            for i in items]
    seeded = [r for r in map(json.loads, SEEDED.read_text(encoding="utf-8").splitlines()) if r and r["kept"]]
    for r in seeded:
        it = by_id[r["source_id"]]
        a = {**answer_of(runs[it["id"]]["answer"]), **{k: r["answer"][k] for k in ("text", "claims")}}
        jobs.append({"key": r["key"], "kind": "seeded", "source_id": it["id"], "transform": r["transform"],
                     "input": VerifierInput(it["question"], it["company"], a, retrieved(it["id"]))})
    random.Random(ORDER_SEED).shuffle(jobs)
    return by_id, jobs


def run(prompt_version: str, use_model: bool) -> dict[str, Any]:
    items, gold, runs = load_dev()
    by_id, jobs = inputs(items, runs)
    store = FactStore()
    gate = ModelGate() if use_model else None
    spent_before = gate.spent() if gate else 0.0
    rows = []
    for mode, quote_rule in (("full", True), ("quote_off", False)):
        v = Verifier(store, gate, quote_rule=quote_rule, prompt_version=prompt_version)
        for j in jobs:
            verdict = v.verify(j["input"])
            rows.append({k: j[k] for k in ("key", "kind", "source_id", "transform")} | {"mode": mode} | verdict.to_dict())
    tag = f"dev_{prompt_version}{'' if use_model else '_nomodel'}"
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / f"{tag}.jsonl").write_text("".join(json.dumps(r) + "\n" for r in rows), encoding="utf-8")
    summary = summarise(rows, items, gold, runs, by_id, prompt_version, use_model)
    summary["verifier_spend_usd"] = round((gate.spent() if gate else 0.0) - spent_before, 6)
    summary["verifier_spend_phase4_total_usd"] = round(gate.spent(), 6) if gate else 0.0
    (OUT_DIR / f"{tag}_summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    write_sheet(rows, jobs, by_id, gold, SHEET_DIR / f"verifier_{tag}.xlsx")
    return summary


def _rate(k: int, n: int) -> dict[str, Any]:
    if not n:
        return {"k": 0, "n": 0, "rate": None, "wilson95": None}
    lo, hi = wilson(k, n)
    return {"k": k, "n": n, "rate": round(k / n, 4), "wilson95": [round(lo, 4), round(hi, 4)]}


def summarise(rows: list[dict], items: list[dict], gold: dict, runs: dict, by_id: dict, prompt_version: str,
              use_model: bool) -> dict[str, Any]:
    scores = {s.id: s for s in (score_answer(i, gold[i["id"]], runs[i["id"]]["answer"],
                                             judge=load_judge().get(f"{i['id']}:baseline")) for i in items)}
    out: dict[str, Any] = {"prompt_version": prompt_version, "prompt_sha256": prompt_sha256(prompt_version),
                           "model_used": use_model, "descriptive_only": True, "modes": {}}
    for mode in ("full", "quote_off"):
        mr = [r for r in rows if r["mode"] == mode]
        seeded = [r for r in mr if r["kind"] == "seeded"]
        natural = [r for r in mr if r["kind"] == "natural"]
        per = {t: _rate(sum(r["decision"] == "block" for r in seeded if r["transform"] == t),
                        sum(1 for r in seeded if r["transform"] == t)) for t in sorted({r["transform"] for r in seeded})}
        macro = [per[t]["rate"] for t in MACRO if per.get(t, {}).get("n")]
        correct = [r for r in natural if scores[r["source_id"]].outcome == "correct"]
        returned_correct = [r for r in correct if scores[r["source_id"]].returned]
        controls = [r for r in natural if by_id[r["source_id"]]["attack_type"] == "control"]
        treat = {}
        for r in natural:
            i = by_id[r["source_id"]]
            judge = load_judge().get(f"{i['id']}:baseline")
            treat[i["id"]] = score_answer(i, gold[i["id"]], runs[i["id"]]["answer"], blocked=r["decision"] == "block",
                                          judge=judge)
        reasons = Counter((x["code"], x["basis"]) for r in mr for x in r["reasons"] if x["blocking"])
        out["modes"][mode] = {
            "detection_by_transform": per,
            "detection_pooled": _rate(sum(r["decision"] == "block" for r in seeded), len(seeded)),
            "detection_macro_T1_T3_T4_T5_T6": round(statistics.mean(macro), 4) if macro else None,
            "seeded_flagged_not_blocked": sum(r["decision"] == "flag" for r in seeded),
            "false_block_rate_correct": _rate(sum(r["decision"] == "block" for r in correct), len(correct)),
            "false_block_rate_correct_returned": _rate(sum(r["decision"] == "block" for r in returned_correct),
                                                       len(returned_correct)),
            "natural_blocks": sorted(r["source_id"] for r in natural if r["decision"] == "block"),
            "natural_flags": sorted(r["source_id"] for r in natural if r["decision"] == "flag"),
            "control_abstentions_added": sum(1 for r in controls if r["decision"] == "block"
                                             and not scores[r["source_id"]].abstained),
            "treatment_outcomes": dict(Counter(s.outcome for s in treat.values())),
            "baseline_outcomes": dict(Counter(s.outcome for s in scores.values())),
            "blocks_by_reason_and_basis": {f"{EVAL_LABEL.get(c, c)} [{b}]": n for (c, b), n in sorted(reasons.items())},
            "model_calls": sum(r["model_calls"] for r in mr),
            "model_calls_per_answer": round(sum(r["model_calls"] for r in mr) / len(mr), 4),
            "rule_s_mean": round(statistics.mean(r["rule_s"] for r in mr), 4),
            "model_s_mean_when_called": round(statistics.mean([r["model_s"] for r in mr if r["model_calls"]]), 3)
            if any(r["model_calls"] for r in mr) else None,
        }
    return out


def write_sheet(rows: list[dict], jobs: list[dict], by_id: dict, gold: dict, path: Path) -> None:
    """Every block and every missed seeded error, both modes, for the owner's review."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    answers = {j["key"]: j["input"].answer["text"] for j in jobs}
    wb = Workbook()
    ws = wb.active
    ws.title = "blocks and misses"
    ws.append(["Mode", "Row kind", "Key", "Source item", "Type", "Subtype", "Transform", "Question", "Answer seen by verifier",
               "Gold value", "Decision", "Reasons (code [basis]: detail)", "Why listed", "Your check", "Your note"])
    for r in sorted(rows, key=lambda r: (r["mode"], r["kind"], r["source_id"], r["transform"] or "")):
        missed = r["kind"] == "seeded" and r["decision"] != "block"
        if r["decision"] != "block" and not missed:
            continue
        i = by_id[r["source_id"]]
        reasons = "\n".join(f"{EVAL_LABEL.get(x['code'], x['code'])} [{x['basis']}]: {x['detail']}" for x in r["reasons"])
        ws.append([r["mode"], r["kind"], r["key"], i["id"], i["attack_type"], i["subtype"], r["transform"] or "",
                   i["question"], answers[r["key"]], gold[i["id"]]["gold"].get("printed_as"), r["decision"], reasons,
                   "missed seeded error" if missed else "block", "", ""])
    widths = [9, 8, 16, 9, 14, 16, 8, 40, 70, 12, 8, 60, 16, 12, 30]
    for n, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, n).column_letter].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "C2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def force_restated(prompt_version: str) -> dict[str, Any]:
    """The model path forced on every untouched dev `restated` answer (D-059): one call each, rules
    skipped. None should be blocked; the model's reasons are kept for review."""
    items, _, runs = load_dev()
    _, jobs = inputs(items, runs)
    restated = {i["id"] for i in items if i["attack_type"] == "restated"}
    gate = ModelGate()
    spent_before = gate.spent()
    v = Verifier(FactStore(), gate, prompt_version=prompt_version)
    rows = []
    for j in sorted((j for j in jobs if j["kind"] == "natural" and j["source_id"] in restated), key=lambda j: j["key"]):
        verdict = v.model_only(j["input"])
        rows.append({"source_id": j["source_id"], "decision": verdict.decision, "model_calls": verdict.model_calls,
                     "reasons": [asdict(r) for r in verdict.reasons]})
    out = {"prompt_version": prompt_version, "prompt_sha256": prompt_sha256(prompt_version), "rows": rows,
           "blocked": sum(r["decision"] == "block" for r in rows),
           "verifier_spend_usd": round(gate.spent() - spent_before, 6)}
    path = OUT_DIR / f"dev_forced_model_restated_{prompt_version}.json"
    path.write_text(json.dumps(out, indent=2) + "\n", encoding="utf-8")
    return out


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--prompt", default=PROMPT_VERSION)
    parser.add_argument("--no-model", action="store_true")
    parser.add_argument("--force-restated", action="store_true", help="force the model path on dev restated answers")
    args = parser.parse_args()
    if args.force_restated:
        print(json.dumps(force_restated(args.prompt), indent=2))
        return
    summary = run(args.prompt, not args.no_model)
    for mode, m in summary["modes"].items():
        print(mode, json.dumps({k: m[k] for k in ("detection_by_transform", "detection_macro_T1_T3_T4_T5_T6",
                                                  "false_block_rate_correct", "natural_blocks", "model_calls")}))
    print("spend", summary["verifier_spend_usd"], "phase4 verifier total", summary["verifier_spend_phase4_total_usd"])


if __name__ == "__main__":
    main()
