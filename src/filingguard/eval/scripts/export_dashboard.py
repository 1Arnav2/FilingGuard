"""Phase 6b: per-item dashboard export (D-075). No model calls; no existing results file changed.

The dashboard API reads only `results/`, and only `eval/` reads gold. This script joins, for the 210
test items, the frozen answers, scores, verdicts, judge decisions, gold, the D-071 corrections and the
owner's hand-check into one file the API can serve as is. The corrected basis is applied with the
Phase 5 summary's own `overlay` (imported, not re-stated), so every per-item outcome adds up to
`results/summary.json`. Writes `results/dashboard/items_v1.jsonl` and `results/dashboard/meta_v1.json`.

Usage: python src/filingguard/eval/scripts/export_dashboard.py
"""

import hashlib
import json
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))
from summary_test import overlay  # noqa: E402  (the Phase 5 summary's D-071 overlay)

from filingguard.eval.dev_report import GOLD, _gold_text  # noqa: E402
from filingguard.eval.dev_run import ITEMS, load_items  # noqa: E402
from filingguard.eval.score import Score  # noqa: E402

R = Path("results")
RUNS = R / "runs/test_target_v1.jsonl"
BASE = R / "test/baseline_scores_v1.jsonl"
TREAT = R / "test/treatment_scores_v1.jsonl"
CORRECTIONS = R / "test/baseline_corrections_v1.json"
VERDICTS = R / "test/verifier_v1.jsonl"
JUDGE = R / "judge/test_judge_v1.jsonl"
SHEET = Path("review/test_handcheck_v1.xlsx")
DECISIONS = Path("docs/DECISIONS_LOG.md")
OUT_ITEMS = R / "dashboard/items_v1.jsonl"
OUT_META = R / "dashboard/meta_v1.json"
TAGS = ("target-v1", "testset-v1", "verifier-v1")
REPO = "https://github.com/1Arnav2/FilingGuard"
JUDGE_FIELDS = ("rejects_premise", "confirms_change", "gives_reasons", "reasons_in_filing", "says_no_reasons_given",
                "reason")
SCORE_FIELDS = ("outcome", "label", "notes", "abstained", "scale_not_followed", "unparseable", "injection",
                "answer_followed")


def jl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def handcheck_rows() -> dict[str, dict[str, Any]]:
    """The owner's verdict and note for each natural row of the test hand-check sheet."""
    from openpyxl import load_workbook

    ws = load_workbook(SHEET, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    col = {name: i for i, name in enumerate(rows[0])}
    out = {}
    for r in rows[1:]:
        if r[col["Row kind"]] != "natural":
            continue
        out[r[col["Item"]]] = {"why_listed": r[col["Why listed"]], "verdict": r[col["Your verdict"]],
                               "note": r[col["Your note"]], "pre_filled": r[col["Already decided (D-071)"]] or None}
    return out


def answer_view(answer: dict[str, Any]) -> dict[str, Any]:
    """The answer as stored, except that an unparseable runaway reply is not reproduced."""
    if answer.get("parse_error"):
        return {"text": None, "abstained": answer["abstained"], "claims": [], "citations": [],
                "parse_error": answer["parse_error"], "unparseable_chars": len(answer["text"] or "")}
    return {"text": answer["text"], "abstained": answer["abstained"], "claims": answer["claims"],
            "citations": answer["citations"], "parse_error": None, "unparseable_chars": None}


def build() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Every test item's dashboard row, and the method facts with input hashes."""
    items = load_items("test", allow_test=True)
    by_id = {i["id"]: i for i in items}
    gold = {r["id"]: r for r in jl(GOLD) if r["id"] in by_id}
    runs = {r["id"]: r for r in jl(RUNS)}
    base = {d["id"]: Score(**d) for d in jl(BASE)}
    treat = {d["id"]: Score(**d) for d in jl(TREAT)}
    corr = json.loads(CORRECTIONS.read_text(encoding="utf-8"))
    fixes = {c["item"]: c for c in corr["corrections"]}
    kept = {c["item"]: c for c in corr["not_corrected"]}
    scores = {"as_scored": {"baseline": base, "treatment": treat},
              "corrected": {"baseline": overlay(base, fixes), "treatment": overlay(treat, fixes)}}
    verdicts: dict[str, dict[str, dict]] = {"full": {}, "quote_off": {}}
    for v in jl(VERDICTS):
        if v["kind"] == "natural":
            verdicts[v["mode"]][v["source_id"]] = v
    judge = {(r["id"], r["condition"]): r for r in jl(JUDGE)}
    hand = handcheck_rows()

    rows = []
    for i in sorted(by_id):
        it, g, run = by_id[i], gold[i], runs[i]
        if i in fixes:
            c = fixes[i]
            correction = {"status": "corrected", "as_scored": c["as_scored"], "corrected": c["corrected"],
                          "pattern": c["pattern"], "evidence": c["evidence"], "reason": c["reason"],
                          "approval": c["approval"]}
        elif i in kept:
            correction = {"status": "kept as scored", "as_scored": kept[i]["as_scored"], "corrected": None,
                          "pattern": None, "evidence": None, "reason": kept[i]["reason"], "approval": None}
        else:
            correction = None
        rows.append({
            "id": i, "company": it["company"], "attack_type": it["attack_type"], "subtype": it["subtype"],
            "question": it["question"],
            "gold": {"expected_behavior": g["expected_behavior"], "text": _gold_text(it, g), **g["gold"]},
            "answer": answer_view(run["answer"]),
            "retrieved": run["retrieved"],
            "figures": base[i].figures,
            # the judge read the one answer once (baseline); a passed treatment answer keeps that decision
            "judge": {k: judge[(i, "baseline")].get(k) for k in JUDGE_FIELDS} if (i, "baseline") in judge else None,
            "outcomes": {b: {c: {k: getattr(scores[b][c][i], k) for k in SCORE_FIELDS} for c in ("baseline", "treatment")}
                         for b in ("corrected", "as_scored")},
            "verdict": {m: {"decision": verdicts[m][i]["decision"], "reasons": verdicts[m][i]["reasons"]}
                        for m in ("full", "quote_off")},
            "correction": correction,
            "handcheck": hand.get(i),
        })

    all_items = jl(ITEMS)
    split = dict(Counter(r["split"] for r in all_items))
    d056 = re.search(r"### D-056 — (.+)\n\*\*Date:\*\* (\d{4}-\d{2}-\d{2})", DECISIONS.read_text(encoding="utf-8"))
    tags = {}
    for t in TAGS:
        commit = subprocess.run(["git", "rev-list", "-n", "1", t], capture_output=True, text=True, check=True).stdout.strip()
        tags[t] = {"commit": commit, "url": f"{REPO}/tree/{commit}"}
    inputs = [RUNS, BASE, TREAT, CORRECTIONS, VERDICTS, JUDGE, GOLD, ITEMS, SHEET]
    meta = {
        "generated_by": "src/filingguard/eval/scripts/export_dashboard.py",
        "items_file": OUT_ITEMS.as_posix(),
        "split": {"dev": split["dev"], "test": split["test"], "total": len(all_items)},
        "attack_types_test": dict(sorted(Counter(r["attack_type"] for r in items).items())),
        "companies": len({r["company"] for r in all_items}),
        "freeze_tags": tags,
        "repo": REPO,
        "pass_bars_set": {"decision": "D-056", "title": d056.group(1), "date": d056.group(2), "where": "EVAL §6"},
        "corrections": {"decision": corr["decision"], "rule": corr["rule"], "corrected": len(fixes),
                        "kept_as_scored": len(kept)},
        "handcheck": {"sheet": SHEET.as_posix(), "natural_rows": len(hand),
                      "with_note": sum(bool(h["note"]) for h in hand.values())},
        "inputs_sha256": {p.as_posix(): sha256(p) for p in inputs},
    }
    return rows, meta


def main() -> None:
    rows, meta = build()
    OUT_ITEMS.parent.mkdir(parents=True, exist_ok=True)
    OUT_ITEMS.write_bytes("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows).encode("utf-8"))
    OUT_META.write_bytes((json.dumps(meta, indent=2, ensure_ascii=False) + "\n").encode("utf-8"))
    print(json.dumps({"items": len(rows), "split": meta["split"], "tags": {k: v["commit"][:7] for k, v in
                                                                            meta["freeze_tags"].items()},
                      "pass_bars_set": meta["pass_bars_set"], "handcheck": meta["handcheck"]}, indent=1))


if __name__ == "__main__":
    main()
