"""Phase 5 step 9 (prepare only): the owner's hand-check sheet (EVAL §10 4a, D-062). No model calls.

Rows, deduplicated:
- 30 natural test rows drawn by seed `20261006-handcheck`, stratified by type (allocation in
  proportion to each type's size, largest remainder, ties by type name), each with a baseline and a
  treatment verdict column (60 decisions);
- every natural row not Correct at baseline as scored;
- every natural answer the verifier blocked or flagged (full mode, as in treatment);
- every seeded copy the verifier did not block, in either mode.
Rows the owner already decided in D-071 are marked pre-filled. Agreement on the random 30 is reported
separately from the other rows.

Usage: python src/filingguard/eval/scripts/handcheck_sheet_test.py
"""

import json
import math
import random
from collections import Counter
from pathlib import Path

from filingguard.eval.dev_report import GOLD, _gold_text
from filingguard.eval.dev_run import load_items

SEED = "20261006-handcheck"
N_RANDOM = 30
RUNS = Path("results/runs/test_target_v1.jsonl")
BASE = Path("results/test/baseline_scores_v1.jsonl")
TREAT = Path("results/test/treatment_scores_v1.jsonl")
VERDICTS = Path("results/test/verifier_v1.jsonl")
SEEDED = Path("results/test/seeded_v1.jsonl")
CORRECTIONS = Path("results/test/baseline_corrections_v1.json")
JUDGE = Path("results/judge/test_judge_v1.jsonl")
SHEET = Path("review/test_handcheck_v1.xlsx")


def stratified_sample(items: list[dict]) -> list[str]:
    by_type: dict[str, list[str]] = {}
    for i in sorted(items, key=lambda x: x["id"]):
        by_type.setdefault(i["attack_type"], []).append(i["id"])
    quota = {t: N_RANDOM * len(ids) / len(items) for t, ids in by_type.items()}
    alloc = {t: math.floor(q) for t, q in quota.items()}
    for t in sorted(quota, key=lambda t: (-(quota[t] - alloc[t]), t))[:N_RANDOM - sum(alloc.values())]:
        alloc[t] += 1
    rng = random.Random(SEED)
    return sorted(i for t in sorted(by_type) for i in rng.sample(by_type[t], alloc[t]))


def jl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def figures_text(s: dict) -> str:
    return "\n".join(f"{f['value']:,.6g} [{f['scale'] or '-'}] metric={f['metric']} asked={f['asked']} "
                     f"period={f['period']} ok={f['period_ok']}" for f in s["figures"]) or "(none)"


def build() -> tuple[list[list], dict]:
    items = load_items("test", allow_test=True)
    by_id = {i["id"]: i for i in items}
    gold = {r["id"]: r for r in jl(GOLD) if r["id"] in by_id}
    runs = {r["id"]: r for r in jl(RUNS)}
    base, treat = {r["id"]: r for r in jl(BASE)}, {r["id"]: r for r in jl(TREAT)}
    verdicts = jl(VERDICTS)
    seeded = {r["key"]: r for r in jl(SEEDED) if r["kept"]}
    judge = {r["id"]: r for r in jl(JUDGE)}
    corr = json.loads(CORRECTIONS.read_text(encoding="utf-8"))
    decided = {c["item"]: f"D-071: corrected to Correct ({c['pattern']})" for c in corr["corrections"]}
    decided |= {c["item"]: f"D-071: kept as scored ({c['reason']})" for c in corr["not_corrected"]}

    natural_full = {v["source_id"]: v for v in verdicts if v["kind"] == "natural" and v["mode"] == "full"}
    random30 = stratified_sample(items)
    why: dict[str, list[str]] = {}
    for i in random30:
        why.setdefault(i, []).append("random 30")
    for i, s in base.items():
        if s["outcome"] != "correct":
            why.setdefault(i, []).append("baseline not Correct")
    for i, v in natural_full.items():
        if v["decision"] != "pass":
            why.setdefault(i, []).append(f"verifier {v['decision']}")
    missed: dict[str, list[str]] = {}
    for v in verdicts:
        if v["kind"] == "seeded" and v["decision"] != "block":
            missed.setdefault(v["key"], []).append(f"{v['mode']}: {v['decision']}")

    rows = []
    for i in sorted(why):
        it, s, t, v = by_id[i], base[i], treat[i], natural_full[i]
        j = judge.get(i)
        jtext = "" if not j else f"rejects={j.get('rejects_premise')} confirms={j.get('confirms_change')}\n{j.get('reason')}"
        reasons = "\n".join(f"{x['code']} [{x['basis']}]{' (blocking)' if x['blocking'] else ''}: {x['detail']}"
                            for x in v["reasons"])
        pre = decided.get(i, "")
        rows.append(["natural", "; ".join(why[i]), i, "", it["attack_type"], it["subtype"], it["company"], it["question"],
                     _gold_text(it, gold[i]), runs[i]["answer"]["text"] if not runs[i]["answer"].get("parse_error")
                     else "[unparseable: runaway reply cut off at the output limit]", figures_text(s), jtext,
                     f"{s['outcome']} {s['label'] or ''}".strip(), f"{v['decision']}\n{reasons}".strip(),
                     f"{t['outcome']} {t['label'] or ''}".strip(), pre, "pre-filled" if pre else "", "", ""])
    for k in sorted(missed, key=lambda k: (seeded[k]["source_id"], seeded[k]["transform"])):
        r = seeded[k]
        it = by_id[r["source_id"]]
        reasons = "\n".join(f"{v['mode']}: {v['decision']} "
                            + "; ".join(f"{x['code']}: {x['detail']}" for x in v["reasons"])
                            for v in verdicts if v["key"] == k)
        rows.append(["seeded copy", "missed seeded error (" + ", ".join(missed[k]) + ")", r["source_id"], k,
                     it["attack_type"], it["subtype"], it["company"], it["question"],
                     _gold_text(it, gold[r["source_id"]]) + f"\nTRANSFORM: {r['transform']}", r["answer"]["text"], "", "",
                     f"gate: {r['gate_outcome']} {r['gate_label']}", reasons, "", "", "", "", ""])
    info = {"natural_rows": len(why), "seeded_rows": len(missed), "rows": len(rows), "random30": random30,
            "random30_by_type": dict(Counter(by_id[i]["attack_type"] for i in random30)),
            "prefilled": sorted(i for i in why if i in decided),
            "reasons": dict(Counter(x.split(" (")[0] for w in why.values() for x in w))}
    return rows, info


def main() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    rows, info = build()
    wb = Workbook()
    ws = wb.active
    ws.title = "test hand-check"
    ws.append(["Row kind", "Why listed", "Item", "Seeded key", "Type", "Subtype", "Company", "Question", "Gold",
               "Answer", "Figures read (scorer)", "Judge", "Baseline (as scored)", "Verifier (full mode)",
               "Treatment (as scored)", "Already decided (D-071)", "Pre-filled", "Your verdict", "Your note"])
    for r in rows:
        ws.append(r)
    widths = [11, 26, 9, 16, 14, 18, 8, 40, 34, 70, 44, 30, 16, 44, 16, 34, 10, 22, 30]
    for n, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, n).column_letter].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "D2"
    SHEET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SHEET)
    print(json.dumps(info, indent=2))


if __name__ == "__main__":
    main()
